"""md_to_docx — 送审 md → 公文风 .docx.

pandoc 负责 md→docx 结构转换，python-docx 后处理套用中文公文格式：
  正文   宋体 小四(12pt) 1.5倍行距 首行缩进2字符
  标题1  黑体 16pt 加粗
  标题2  黑体 14pt 加粗
  标题3  黑体 12pt 加粗
  表格   宋体 10.5pt，表头加粗居中浅底
  页边距 上下2.54 左右3.18cm（财评附件通用）
"""
from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BODY_CN = "宋体"
HEAD_CN = "黑体"
BODY_PT = 12          # 小四
TABLE_PT = 10.5       # 五号
_TODO_MARKERS = ("待核价", "待核", "待补", "待查证", "待确认", "待技术",
                 "待商务", "待填", "____", "【待", "（待", "型号/单价待核")


def _set_run_font(run, cn_font: str, size_pt: float, bold=False, color=None):
    run.font.name = cn_font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    # 中文字形必须设 eastAsia
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn_font)
    rfonts.set(qn("w:ascii"), cn_font)
    rfonts.set(qn("w:hAnsi"), cn_font)


def _style_doc(docx_path: Path) -> None:
    doc = Document(str(docx_path))

    # 页边距
    for sec in doc.sections:
        sec.top_margin = Cm(2.54)
        sec.bottom_margin = Cm(2.54)
        sec.left_margin = Cm(3.18)
        sec.right_margin = Cm(3.18)

    for p in doc.paragraphs:
        style = (p.style.name or "").lower()
        text = p.text.strip()
        if not text:
            continue
        if style.startswith("title") or style == "heading 1":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if style.startswith("title") else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                _set_run_font(r, HEAD_CN, 16, bold=True, color=(0x1F, 0x3B, 0x73))
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(12)
        elif style == "heading 2":
            for r in p.runs:
                _set_run_font(r, HEAD_CN, 14, bold=True)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(6)
        elif style == "heading 3":
            for r in p.runs:
                _set_run_font(r, HEAD_CN, 12, bold=True)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(4)
        else:
            todo = any(m in text for m in _TODO_MARKERS)
            for r in p.runs:
                if todo and any(m in (r.text or "") for m in _TODO_MARKERS):
                    _set_run_font(r, BODY_CN, BODY_PT, color=(0x1F, 0x66, 0xE5))
                    r.font.italic = True
                else:
                    _set_run_font(r, BODY_CN, BODY_PT)
            pf = p.paragraph_format
            pf.line_spacing = 1.5
            pf.first_line_indent = Pt(BODY_PT * 2)   # 首行缩进2字符
            pf.space_after = Pt(2)

    def _force_grid(tbl):
        """OOXML 强制全网格细线（不依赖样式名是否存在）。"""
        tblPr = tbl._tbl.tblPr
        borders = tblPr.makeelement(qn("w:tblBorders"), {})
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = borders.makeelement(qn(f"w:{edge}"), {
                qn("w:val"): "single", qn("w:sz"): "4",
                qn("w:space"): "0", qn("w:color"): "BFBFBF",
            })
            borders.append(el)
        old = tblPr.find(qn("w:tblBorders"))
        if old is not None:
            tblPr.remove(old)
        tblPr.append(borders)

    # 表格
    for tbl in doc.tables:
        try:
            tbl.style = "Table Grid"
        except Exception:
            pass
        _force_grid(tbl)
        for i, row in enumerate(tbl.rows):
            for cell in row.cells:
                cell_todo = i > 0 and any(m in cell.text for m in _TODO_MARKERS)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in (p.runs or [p.add_run("")]):
                        if cell_todo and any(m in (r.text or "") for m in _TODO_MARKERS):
                            _set_run_font(r, BODY_CN, TABLE_PT, color=(0x1F, 0x66, 0xE5))
                            r.font.italic = True
                        else:
                            _set_run_font(r, BODY_CN, TABLE_PT, bold=(i == 0))
                if i == 0:
                    # 表头浅底
                    tcPr = cell._tc.get_or_add_tcPr()
                    shd = tcPr.makeelement(qn("w:shd"),
                                           {qn("w:val"): "clear",
                                            qn("w:color"): "auto",
                                            qn("w:fill"): "DCE6F1"})
                    tcPr.append(shd)

    doc.save(str(docx_path))


def convert(md_path, out_dir) -> Path:
    md_path = Path(md_path)
    out_dir = Path(out_dir)
    docx_path = out_dir / (md_path.stem + ".docx")
    subprocess.run(
        ["pandoc", str(md_path), "-o", str(docx_path),
         "--from", "gfm", "--to", "docx"],
        check=True, capture_output=True,
    )
    _style_doc(docx_path)
    return docx_path


if __name__ == "__main__":  # pragma: no cover
    import sys
    print(convert(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "."))
