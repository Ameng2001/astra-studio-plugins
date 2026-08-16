"""smoke — 端到端冒烟矩阵：拿**真实项目**跑全链路，每个标准包各跑一遍。

## 为什么需要它

`test_guards.py` 用合成数据测单元。它有价值，但有个盲区：
**这个项目的缺陷几乎全部只在真实输入上显形。** 一轮改造下来修的 11 个缺陷里——

  8 个  在真实数据上跑命令时炸出来（¥40,000 实施费未列进清单、明细用了
        另一个没有交付方案的引擎实例、复算式漏 ROUND、build_fp_table
        反推 FP 归零、run_review 报 0/100、年度金额列认不出…）
  2 个  手工查出来（汇总表靠碰巧被识别、4 条 rationale 与 type 互相矛盾）
  1 个  别人问「换个省行不行」才发现（广东 emit_whatif KeyError）
  0 个  被 test_guards 抓到

最后一条最说明问题：广东那个 `KeyError: 'dev_first'` 早就躺在那儿了，
而单元测试全绿。**「在一个输入上通过」和「通用」是两回事。**

所以这个矩阵有两条硬要求：

  1. **至少两个标准包** —— 只跑山东等于什么都没验。跨包的形状差异
     （广东 `dev_category.values: {}`、应用类型别名、无其他费用科目）
     正是最容易漏的地方。
  2. **断言具体数字，不只断言 exit 0** —— 那些静默出 0 的缺陷全都是
     退出码 0 的。「跑通了」不等于「算对了」。

## 用法

    BOM_PROJECT_ROOT=<项目目录> python3 tests/smoke.py [--quick]

`--quick` 跳过 review 链路（最慢的一段），只验生成侧。
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
ROOT = Path(os.environ.get("BOM_PROJECT_ROOT", "")).resolve() if os.environ.get(
    "BOM_PROJECT_ROOT") else None

FAILURES: list[str] = []
SKIPS: list[str] = []
DELIVERY: list[str] = []


def check(label: str, got, want) -> bool:
    if got == want:
        print(f"    ✓ {label}")
        return True
    FAILURES.append(f"{label}: 实得 {got!r}，期望 {want!r}")
    print(f"    ✗ {label}: 实得 {got!r}，期望 {want!r}")
    return False


def close(label: str, got: float, want: float, tol: float = 0.01) -> bool:
    if abs(got - want) <= tol:
        print(f"    ✓ {label}　{got:,.2f}")
        return True
    FAILURES.append(f"{label}: 实得 {got:,.2f}，期望 {want:,.2f}（容差 {tol}）")
    print(f"    ✗ {label}: 实得 {got:,.2f}，期望 {want:,.2f}")
    return False


def run(argv: list[str], cwd: Path, label: str) -> subprocess.CompletedProcess:
    env = {**os.environ, "PYTHONPATH": str(SCRIPTS)}
    p = subprocess.run([sys.executable, *argv], cwd=cwd, env=env,
                       capture_output=True, text=True)
    if p.returncode != 0:
        FAILURES.append(f"{label} 退出码 {p.returncode}\n"
                        + "\n".join(p.stderr.strip().splitlines()[-12:]))
        print(f"    ✗ {label} 退出码 {p.returncode}")
        print("      " + "\n      ".join(p.stderr.strip().splitlines()[-8:]))
    else:
        print(f"    ✓ {label}")
    return p


# ---- LibreOffice 重算 ---------------------------------------------------

def _soffice() -> str | None:
    for c in ("soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if shutil.which(c) or Path(c).exists():
            return c
    return None


def recalc(paths: list[Path], out: Path) -> dict[str, Path] | None:
    """用 LibreOffice 重算公式。

    **这是活公式唯一的真实验证。** 生成时的 `gov_sheet.formula()` 自验只证明
    「Python 按同一算式算的结果对」，证明不了「Excel 解析这个公式串没问题」——
    引用写错表名、区间越界、VLOOKUP 查不到，都要真算一遍才知道。

    LibreOffice 没装时**大声跳过**，不静默通过 —— 否则等于悄悄少验一层。
    """
    so = _soffice()
    if not so:
        SKIPS.append("LibreOffice 未安装 —— **活公式未做真实重算**，"
                     "只验了生成时的 Python 自验")
        print("    ⚠ 跳过公式重算：LibreOffice 未安装")
        return None
    out.mkdir(parents=True, exist_ok=True)
    subprocess.run([so, "--headless", "--calc", "--convert-to", "xlsx",
                    "--outdir", str(out), *[str(p) for p in paths]],
                   capture_output=True, text=True, timeout=600)
    got = {p.name: out / p.name for p in paths if (out / p.name).exists()}
    if len(got) != len(paths):
        FAILURES.append(f"LibreOffice 重算产出 {len(got)}/{len(paths)} 个文件")
    return got


def total_row(ws, label_prefix: str = "合计") -> dict[str, object]:
    """取合计行，按表头列名返回。"""
    hdr_row = next((r for r in range(1, 6)
                    if ws.cell(r, 1).value == "序号"), 3)
    h = {ws.cell(hdr_row, c).value: c for c in range(1, ws.max_column + 1)}
    for r in range(hdr_row + 1, ws.max_row + 1):
        for c in range(1, 4):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.startswith(label_prefix):
                return {k: ws.cell(r, i).value for k, i in h.items() if k}
    return {}


# ---- 一个标准包跑一轮 ---------------------------------------------------

def find_delivery(root: Path) -> list[str]:
    """项目里有交付方案就带上 —— 否则 Z03/Z04 根本不生成，
    `emit_ops_list` / `emit_tco` 一行都跑不到。

    **覆盖范围要跟着项目实际有什么走**，不是硬写「必须 5 册」：
    没有交付方案时 3 册才是对的。首版硬断言 5 册，两个包都红，
    而工具没有任何问题 —— 断言本身错了。
    """
    modes = root / "delivery-modes/modes.yaml"
    if not modes.exists():
        return []
    # 优先 deal.yaml（第三层决策的单一来源）；退回 delivery-plan.yaml 兼容旧项目。
    # 找不到就只跑生成侧 —— 但覆盖缺口要 SKIPS 里说出来，不静默少验。
    deal = next(iter(sorted((root / "deals").glob("*/deal.yaml"))), None)
    if deal is not None:
        return ["--deal", str(deal)]
    plan = next(iter(sorted((root / "deals").glob("*/delivery-plan.yaml"))), None)
    if plan is None:
        return []
    argv = ["--modes", str(modes), "--delivery-plan", str(plan)]
    ic = root / "delivery-modes/internal-cost-model.yaml"
    if ic.exists():
        argv += ["--internal-cost", str(ic)]
    sc = sorted((root / "delivery-modes/scenarios").glob("*.yaml"))
    if sc:
        argv += ["--scenarios", *[str(x) for x in sc]]
    return argv


def one_pack(pack_dir: Path, work: Path, quick: bool) -> None:
    import openpyxl

    pid = pack_dir.name
    print(f"\n{'=' * 72}\n■ {pid}\n{'=' * 72}")

    # --- 1. 第二层：区域基准 ---
    print("  [1/5] baseline_build")
    # --out 指到临时目录：**冒烟不改项目产物**。
    # 首版没加，直接覆盖了 baselines/ 下的真实基准；而且 glob 出多个历史
    # 版本目录（@bom-0.16/0.17/0.18）时 next() 会挑到旧的，
    # 于是第 2 步报「基准已过期」—— 那是我的脚本挑错了目录，不是工具的错。
    bl_dir = work / f"baseline-{pid}"
    p = run([str(SCRIPTS / "baseline_build.py"), "--bom", "bom",
             "--pack", str(pack_dir), "--out", str(bl_dir)],
            ROOT, f"{pid} baseline_build")
    if p.returncode != 0:
        return
    if not check(f"{pid} 基准目录存在", bl_dir.is_dir(), True):
        return
    bl = json.loads((bl_dir / "baseline.json").read_text(encoding="utf-8"))
    for f in ("baseline.json", "baseline.lock.json", "citations.md",
              "区域基准清单.xlsx"):
        check(f"{pid} 产物 {f}", (bl_dir / f).exists(), True)

    # --- 2. 第三层：报价 ---
    print("  [2/5] quote_generate")
    deal = work / f"deal-{pid}"
    # --deal 里带了 baseline，但冒烟要用自己刚建的那份（临时目录），
    # 所以 --baseline 显式给 —— CLI 覆盖文件正是为这种场景设计的。
    p = run([str(SCRIPTS / "quote_generate.py"), "--baseline", str(bl_dir),
             "--out", str(deal), "--deal-id", f"冒烟-{pid}",
             "--doc-prefix", "冒烟", "--doc-date", "20260101",
             "--allow-rule-violations", *DELIVERY],
            ROOT, f"{pid} quote_generate")
    if p.returncode != 0:
        return
    out = deal / "out"
    xlsx = sorted(out.glob("*.xlsx"))
    want_n = 5 if DELIVERY else 3          # 无交付方案时 Z03/Z04 不生成
    check(f"{pid} 出 {want_n} 册 xlsx", len(xlsx), want_n)
    for tag in (["Z00", "Z01", "Z02"] + (["Z03", "Z04"] if DELIVERY else [])):
        check(f"{pid} 有 {tag}", any(tag in x.name for x in xlsx), True)
    res = json.loads((out / "costing-result.json").read_text(encoding="utf-8"))
    eng_sw = res["software_dev"]["total"]
    eng_ct = res["totals"]["construction_total"]

    # --- 3. 公式重算：活公式必须算出引擎值 ---
    print("  [3/5] LibreOffice 重算")
    rc = recalc(xlsx, work / f"rc-{pid}")
    if rc:
        for name, path in sorted(rc.items()):
            wb = openpyxl.load_workbook(path, data_only=True)
            for sn in wb.sheetnames:
                t = total_row(wb[sn])
                if not t:
                    continue
                # 采购清单与总报价汇总的合计 = 建设期合计
                for col, want in (("金额（元）", eng_ct),
                                  ("软件开发费用（元）", eng_sw)):
                    if isinstance(t.get(col), (int, float)):
                        close(f"{pid} [{sn}] {col} 合计", t[col], want)
        # 试算表：活公式的合计与引擎值的差必须在已知量级内
        z02 = next((p for n, p in rc.items() if "Z02" in n), None)
        if z02:
            wb = openpyxl.load_workbook(z02, data_only=True)
            check(f"{pid} Z02 有 03_二层试算", "03_二层试算" in wb.sheetnames, True)
            t = total_row(wb["03_二层试算"], "试算合计")
            v = t.get("软件开发费（元）")
            if isinstance(v, (int, float)):
                # 逐条目 xlround 累加 vs 子系统整体相乘，差应 < 0.01%
                close(f"{pid} 试算合计 vs 引擎（≤0.01%）", v, eng_sw,
                      tol=max(1.0, eng_sw * 0.0001))
            else:
                FAILURES.append(f"{pid} 试算合计不是数值：{v!r} —— "
                                f"活公式没被 Excel 算出来")

    if quick:
        print("  [4-5/5] 跳过 review 链路（--quick）")
        return

    # --- 4. 解析 + 优化：不能静默出 0 ---
    print("  [4/5] parse_quote → run_optimize")
    sess = work / f"sess-{pid}"
    sess.mkdir(parents=True, exist_ok=True)
    run([str(SCRIPTS / "parse_quote.py"), str(sess / "quote.json"),
         *[str(x) for x in xlsx]], ROOT, f"{pid} parse_quote")
    std = pack_dir / "standard.parsed.json"
    if std.exists():
        shutil.copy(std, sess / "standard.json")
    p = run([str(SCRIPTS / "run_optimize.py"), str(sess),
             "--standard", str(pack_dir)], ROOT, f"{pid} run_optimize")
    if p.returncode == 0:
        sg = json.loads((sess / "optimize-suggestions.json").read_text(
            encoding="utf-8"))["summary"]
        # **这两条是本矩阵存在的理由**：它们曾经静默为 0 而命令退出 0
        ot = sg["budget_band"]["original_total"]
        check(f"{pid} 原总额被识别（非 0）", ot > 0, True)
        # 键名是 fp_summary 不是 feasibility_fp —— 首版写错了键，
        # `.get(..., {})` 让它静默取 0 然后断言失败。断言写错至少会红；
        # 若这里写的是「>= 0」就会静默通过，那才是最坏的情况。
        check(f"{pid} 功能点被识别（非 0）",
              sg["fp_summary"]["total_fp"] > 0, True)

    # --- 5. 评审：不能把「没得查」印成 0 分 ---
    print("  [5/5] run_review")
    if not (sess / "standard.json").exists():
        SKIPS.append(f"{pid} 无 standard.parsed.json —— 跳过 run_review")
        print(f"    ⚠ 跳过：{pack_dir}/standard.parsed.json 不存在")
        return
    p = run([str(SCRIPTS / "run_review.py"), str(sess),
             "--standard", str(pack_dir)], ROOT, f"{pid} run_review")
    if p.returncode == 0:
        rep = (sess / "review-report.md").read_text(encoding="utf-8")
        # 「无适用行」必须与「全不合格」长得不一样 —— 这曾经是同一个 0/100
        bad = "综合得分**：0/100" in rep or "综合得分**：0/100" in rep
        check(f"{pid} 未把「没得查」印成 0/100", bad, False)


def main() -> int:
    if ROOT is None or not (ROOT / "bom").is_dir():
        print("需要真实项目：BOM_PROJECT_ROOT=<含 bom/ 与 standard-packs/ 的目录>")
        print("\n**这个矩阵不能用合成数据跑。** 它要验的正是真实输入上的形状差异 ——")
        print("合成数据里不会有「广东 dev_category.values 为空」这种东西。")
        return 2
    quick = "--quick" in sys.argv
    global DELIVERY
    DELIVERY = find_delivery(ROOT)
    print("交付方案：" + ("已挂载，Z03/Z04 纳入覆盖"
                          if DELIVERY else "**项目未提供，Z03/Z04 未覆盖**"))
    if not DELIVERY:
        SKIPS.append("项目无 delivery-modes/modes.yaml 或 deals/*/delivery-plan.yaml"
                     " —— emit_ops_list / emit_tco 本轮一行都没跑到")
    packs = sorted(d for d in (ROOT / "standard-packs").iterdir()
                   if d.is_dir() and (d / "pack.yaml").exists()
                   and not d.name.startswith("_"))
    import yaml
    active = []
    for d in packs:
        doc = yaml.safe_load((d / "pack.yaml").read_text(encoding="utf-8"))
        if doc.get("deprecated"):
            # 退役包不参与冒烟：它拒绝加载是**正确行为**，不是缺陷
            SKIPS.append(f"{d.name} 已退役（{doc['deprecated'].get('since')}），"
                         f"不参与冒烟 —— 它拒绝加载是正确行为")
            continue
        active.append(d)
    print(f"项目 {ROOT}")
    print(f"标准包 {len(active)} 个：{[d.name for d in active]}")
    if len(active) < 2:
        FAILURES.append("只有 1 个标准包 —— 单包冒烟等于什么都没验，"
                        "跨包的形状差异才是最容易漏的地方")

    with tempfile.TemporaryDirectory(prefix="fund-smoke-") as td:
        for d in active:
            one_pack(d, Path(td), quick)

    print(f"\n{'=' * 72}")
    for s in SKIPS:
        print(f"⚠ 跳过：{s}")
    if FAILURES:
        print(f"冒烟失败 {len(FAILURES)} 项：")
        for f in FAILURES:
            print(f"  ✗ {f}")
        return 1
    print("冒烟矩阵：全部通过"
          + ("（--quick，未跑 review 链路）" if quick else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
