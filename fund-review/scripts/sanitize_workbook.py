"""sanitize_workbook —— 把内部工作簿脱敏成可外发的送审件。

## 为什么需要它

业务侧的报价工作簿常常把**成本、渠道价、毛利率、供应商**和对外报价放在同一张表
（Z03 场景设备报价就是：成本单价 / 成本总价 / 渠道单价 / 渠道总价 / 市场单价 /
市场总价 六列并排，另有一张「价格测算总览」直接给出毛利率）。

「沿用原表送审」是个合理的业务决定 —— 格式评审方已经看惯了，重排反而增加风险。
但**沿用不等于原样发出去**。这个模块做的就是这件事：保留格式，摘掉不能外发的列。

## 三条纪律

1. **先删公式再删列。** openpyxl 删列不会重写公式引用，`市场总价 =$I3*$W3`
   在删掉左侧四列后仍指向 W 列，而 W 已经是别的东西了 —— 不报错，只是数错了。
   所以一律按**值**另存（`data_only=True`），送审件本来也不该带跨表 VLOOKUP。

2. **删完要验。** 脱敏最危险的失败是「以为删了」。`_verify` 把产物整本重读一遍，
   逐格搜禁用表头与禁用取值；搜到任何一处直接删产物并报错。
   没有这一步的脱敏工具比没有工具更糟 —— 它给人已经安全了的错觉。

3. **删了什么要写在产物里。** 产物第一张表是「脱敏说明」，逐条列出删掉的
   sheet 与列及原因。悄悄删掉一张显示数据质量问题的表，和隐瞒是一回事。

用法：
    python3 sanitize_workbook.py --src <内部.xlsx> --out <送审.xlsx> \\
        --rule <rule.yaml> [--expect-total 23966435]
"""
from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

import openpyxl
import yaml
from openpyxl.styles import Alignment, Font, PatternFill

BODY = "微软雅黑"
_TITLE = Font(name=BODY, size=13, bold=True, color="9C0006")
_HEAD = Font(name=BODY, size=10, bold=True, color="FFFFFF")
_BODY = Font(name=BODY, size=10)
_HEAD_FILL = PatternFill("solid", fgColor="C00000")
_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


class SanitizeError(RuntimeError):
    """脱敏失败。**不降级、不部分成功** —— 宁可没有产物，不可有个看着干净的产物。"""


def sanitize(src: Path, out: Path, rule: dict[str, Any]) -> dict[str, Any]:
    drop_sheets: list[str] = rule.get("drop_sheets") or []
    drop_cols: list[str] = rule.get("drop_columns") or []
    header_rows: dict[str, int] = rule.get("header_rows") or {}
    default_header = rule.get("default_header_row", 2)

    wb = openpyxl.load_workbook(src, data_only=True)   # 值，不带公式
    removed: list[dict[str, Any]] = []

    unknown = set(drop_sheets) - set(wb.sheetnames)
    if unknown:
        raise SanitizeError(
            f"规则要删的 sheet {sorted(unknown)} 在源文件里不存在 —— "
            f"源文件改过名？规则没跟上就等于没删。请先核对再跑。")

    for sn in list(wb.sheetnames):
        if sn in drop_sheets:
            reason = (rule.get("drop_sheet_reasons") or {}).get(sn, "内部数据")
            removed.append({"where": sn, "what": "整张表", "why": reason})
            del wb[sn]
            continue

        ws = wb[sn]
        hr = header_rows.get(sn, default_header)
        hits = [(c, ws.cell(hr, c).value) for c in range(1, ws.max_column + 1)
                if ws.cell(hr, c).value in drop_cols]
        # 从右往左删，否则删了左边的列右边的列号就漂了
        for c, name in sorted(hits, reverse=True):
            ws.delete_cols(c)
            removed.append({"where": sn, "what": f"列「{name}」",
                            "why": (rule.get("drop_column_reasons") or {}).get(
                                name, "内部成本/渠道口径")})

    if not removed:
        raise SanitizeError(
            f"一处都没删 —— 规则与源文件对不上，或源文件已经不含这些列。"
            f"这种情况必须人工确认，不能当成「已经干净了」直接放行。")

    _emit_notice(wb, src, rule, removed)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    _verify(out, drop_cols, drop_sheets, rule.get("forbidden_values") or [])
    return {"removed": removed, "sheets": wb.sheetnames}


def _emit_notice(wb, src: Path, rule: dict, removed: list[dict]) -> None:
    ws = wb.create_sheet("脱敏说明", 0)
    ws["A1"] = "脱敏说明 —— 本件由内部工作簿脱敏生成，以下内容已移除"
    ws["A1"].font = _TITLE
    ws.merge_cells("A1:C1")
    ws["A2"] = f"源文件：{src.name}"
    ws["A2"].font = _BODY
    ws.merge_cells("A2:C2")
    ws["A3"] = (rule.get("notice")
                or "移除的均为内部成本、渠道价、毛利率与供应商信息，"
                   "不影响对外报价金额与设备清单的完整性。")
    ws["A3"].font = _BODY
    ws["A3"].alignment = _WRAP
    ws.merge_cells("A3:C3")

    for i, h in enumerate(("位置", "移除内容", "原因"), start=1):
        c = ws.cell(5, i, h)
        c.font, c.fill, c.alignment = _HEAD, _HEAD_FILL, _WRAP
    for r, item in enumerate(removed, start=6):
        for i, k in enumerate(("where", "what", "why"), start=1):
            c = ws.cell(r, i, item[k])
            c.font, c.alignment = _BODY, _WRAP
    for col, w in (("A", 22), ("B", 26), ("C", 52)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[3].height = 30


def _verify(out: Path, drop_cols: list[str], drop_sheets: list[str],
            forbidden_values: list[str]) -> None:
    """整本重读，逐格搜禁用词。搜到就删产物 —— 不留一个看着干净的文件在磁盘上。"""
    wb = openpyxl.load_workbook(out, data_only=True)
    bad: list[str] = []

    for sn in drop_sheets:
        if sn in wb.sheetnames:
            bad.append(f"sheet「{sn}」仍在")

    needles = [n for n in (drop_cols + forbidden_values) if n]
    for sn in wb.sheetnames:
        if sn == "脱敏说明":       # 说明表本身要写出删了哪些列名，豁免
            continue
        ws = wb[sn]
        for row in ws.iter_rows():
            for cl in row:
                v = cl.value
                if not isinstance(v, str):
                    continue
                for n in needles:
                    if n in v:
                        bad.append(f"[{sn}] {cl.coordinate} 出现「{n}」：{v[:40]}")

    if bad:
        out.unlink(missing_ok=True)
        raise SanitizeError(
            "脱敏产物复检不通过，已删除产物：\n  " + "\n  ".join(bad[:20])
            + "\n脱敏最危险的失败是「以为删了」—— 宁可没有产物。")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--rule", required=True, type=Path)
    ap.add_argument("--expect-total", type=float,
                    help="脱敏后某列的合计期望值，用于确认没删错列")
    ap.add_argument("--total-column", default="市场总价")
    a = ap.parse_args()

    rule = yaml.safe_load(a.rule.read_text(encoding="utf-8"))
    rep = sanitize(a.src, a.out, rule)

    if a.expect_total is not None:
        got = _sum_column(a.out, a.total_column,
                          rule.get("default_header_row", 2))
        if abs(got - a.expect_total) > 0.01:
            a.out.unlink(missing_ok=True)
            raise SanitizeError(
                f"脱敏后「{a.total_column}」合计 {got:,.2f}，期望 {a.expect_total:,.2f}，"
                f"差 {got - a.expect_total:,.2f} —— 删列时把数据删错了，已删除产物")
        print(f"  ✓ 「{a.total_column}」合计 {got:,.2f} 与期望一致")

    print(f"脱敏 → {a.out}")
    print(f"  移除 {len(rep['removed'])} 处，保留 sheet：{rep['sheets']}")
    for x in rep["removed"][:40]:
        print(f"    − {x['where']:<18}{x['what']}")


#: 明细区里的合计行标签。**必须跳过** —— 源表每张园所表末行就是「合计」，
#: 逐行相加会把小计和它的明细各加一遍，结果正好翻倍。
#: 与 gov_sheet 那个「分组小计被 =SUM 重复计入」是同一个形状：不报错，只是数错了。
_TOTAL_LABELS = ("合计", "小计", "总计", "共计")


def _sum_column(path: Path, col_name: str, header_row: int) -> float:
    wb = openpyxl.load_workbook(path, data_only=True)
    total = 0.0
    skipped = 0
    for sn in wb.sheetnames:
        if sn == "脱敏说明":
            continue
        ws = wb[sn]
        idx = next((c for c in range(1, ws.max_column + 1)
                    if ws.cell(header_row, c).value == col_name), None)
        if idx is None:
            continue
        for r in range(header_row + 1, ws.max_row + 1):
            head = [ws.cell(r, c).value for c in range(1, min(4, idx) + 1)]
            if any(isinstance(x, str) and x.strip() in _TOTAL_LABELS
                   for x in head):
                skipped += 1
                continue
            v = ws.cell(r, idx).value
            if isinstance(v, (int, float)):
                total += v
    if skipped:
        print(f"  （校验时跳过 {skipped} 行表内合计行，避免与明细重复计入）")
    return total


if __name__ == "__main__":
    main()
