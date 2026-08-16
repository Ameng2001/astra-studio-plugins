"""bom_effort_wbs —— 给走工作量估算法的条目挂上活动分解（WBS）。

柳州标准 p.11 表1 注1：「系统应按子系统、模块进行细化分层，一般系统分两层，
复杂系统应分三层或更多」。C4–C8 现在是一层（一条一个人天数），挂上活动
分解就是两层。

## 不改共创状态

WBS **不是人月的测算依据** —— 源表里那些分解列是公式
`=ROUND(人天×百分比, 0)`，人天是输入、活动人天是输出。它回答
「126 人天怎么分到 7 个活动」，不回答「126 从哪来」。

所以挂完之后 `★待补工作量依据` 照旧。把 WBS 填进去就标「已判型」，
是用标签盖住缺口 —— 本项目今天已经犯过一次（164 条工作量法条目
曾被标成「已判型」，而它们从来没判过型）。

## 按名称匹配，匹配不上就报错

Z02 的行数与 BOM 逐系统 1:1（知识工程 9 / 数据工程 28 / 专业模型 22 /
智能体 73）。**不做模糊匹配、不按行号对位** —— 行号对位在源表插一行时
会整体错位，而错位的表现是每条都挂上了别人的分解，看起来一切正常。

用法：
    python3 bom_effort_wbs.py --bom <dir> --src <Z02.xlsx> [--apply]
"""
from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path
from typing import Any

import openpyxl
import yaml

#: 各表里「建设内容」那一列的表头名 —— 逐表不同，显式登记
_NAME_COL = {
    "02_知识工程(L1建设)": "建设内容",
    "03_数据建设(L1建设)": "数据集",
    "04_行业专业模型(L2建设)": "模型",
    "05_场景智能体(L2建设)": "智能体",
}


def load_src(src: Path, wbs: dict) -> dict[str, dict[str, dict]]:
    """{系统: {建设内容名: {人天, 数量, 单位, 建设详情}}}"""
    wb = openpyxl.load_workbook(src, data_only=True)
    out: dict[str, dict[str, dict]] = {}
    for system, spec in wbs["templates"].items():
        sheet = spec["sheet"]
        if sheet not in wb.sheetnames:
            raise ValueError(f"源表缺 sheet「{sheet}」")
        ws = wb[sheet]
        h = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1)}
        col = _NAME_COL[sheet]
        rows: dict[str, dict] = {}
        for r in range(3, ws.max_row + 1):
            nm = ws.cell(r, h[col]).value
            if nm is None:
                continue
            nm = str(nm).strip()
            if nm in rows:
                raise ValueError(
                    f"{sheet} 有重名建设内容「{nm}」—— 按名称匹配会挂错，"
                    f"请先在源表消歧")
            detail = next((h[k] for k in ("建设详情", "建设详情（定位）") if k in h),
                          None)
            rows[nm] = {
                "person_days": ws.cell(r, h["人/天"]).value or 0,
                "qty": ws.cell(r, h["数量"]).value,
                "unit": ws.cell(r, h["单位"]).value,
                "detail": (str(ws.cell(r, detail).value or "")[:300]
                           if detail else ""),
                "row": r,
            }
        out[system] = rows
    return out


def attach(bom_dir: Path, src: Path) -> dict[str, Any]:
    wbs = yaml.safe_load((bom_dir / "effort-wbs.yaml").read_text(encoding="utf-8"))
    data = load_src(src, wbs)

    files: dict[Path, Any] = {}
    matched = defaultdict(int)
    unmatched: list[dict] = []
    for f in sorted(bom_dir.glob("items/*.yaml")):
        doc = yaml.safe_load(f.read_text(encoding="utf-8"))
        for it in doc["items"]:
            system = it["path"]["system"]
            if system not in data:
                continue
            row = data[system].get(it["name"].strip())
            if row is None:
                unmatched.append({"id": it["id"], "system": system,
                                  "name": it["name"]})
                continue
            pd = row["person_days"]
            acts = [{"activity": a["name"], "pct": a["pct"],
                     "person_days": round(pd * a["pct"], 1)}
                    for a in wbs["templates"][system]["activities"]]
            eff = dict(it.get("effort") or {})
            eff.update({
                "person_days": pd,
                "man_months": round(pd / 21.75, 4),
                "qty": row["qty"], "unit": row["unit"],
                "wbs": acts,
                "wbs_source": f"{src.name}#{wbs['templates'][system]['sheet']}!"
                              f"r{row['row']}",
                "wbs_note": "活动人天由人天按固定百分比派生（源表为公式），"
                            "是**工作分解**不是测算依据；人天本身的来源仍待补",
            })
            it["effort"] = eff
            matched[system] += 1
        files[f] = doc
    return {"files": files, "matched": dict(matched), "unmatched": unmatched,
            "wbs": wbs}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--bom", required=True, type=Path)
    ap.add_argument("--src", required=True, type=Path)
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()

    rep = attach(a.bom, a.src)
    print(f"挂载活动分解 {'（已写入）' if a.apply else '（dry-run）'}")
    for s, n in sorted(rep["matched"].items()):
        acts = len(rep["wbs"]["templates"][s]["activities"])
        print(f"  {s:<14}{n:>4} 条 × {acts} 个活动")
    if rep["unmatched"]:
        print(f"\n  ⚠ {len(rep['unmatched'])} 条在源表里找不到同名行：")
        for u in rep["unmatched"][:10]:
            print(f"    {u['id']}  {u['system']} / {u['name']}")
        print("  按名称匹配，匹配不上不猜 —— 按行号对位会在源表插一行时整体错位，"
              "而错位的表现是每条都挂上了别人的分解，看起来一切正常")

    if a.apply:
        if rep["unmatched"]:
            raise SystemExit("有条目匹配不上，先对齐名称再 --apply")
        for path, doc in rep["files"].items():
            path.write_text(yaml.safe_dump(doc, allow_unicode=True,
                                           sort_keys=False, width=200),
                            encoding="utf-8")
        (a.bom / "effort-wbs-report.json").write_text(
            json.dumps({"matched": rep["matched"]}, ensure_ascii=False,
                       indent=2), encoding="utf-8")
        print(f"  → {a.bom}")


if __name__ == "__main__":
    main()
