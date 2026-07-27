"""parse_quote — xlsx → quote.json.

Loads one or more quote workbooks and emits a uniform structure used by
optimize/rewrite/review skills.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import openpyxl


# ---- workbook classification ------------------------------------------------
def classify_workbook(name: str, sheet_names: list[str]) -> str:
    n = name.lower()
    sn = " ".join(sheet_names)
    if "大模型" in name or "行业模型" in name or "模型" in name:
        return "llm"
    if "设备" in name or any("园" in s for s in sheet_names):
        return "device"
    if "平台" in name or "软件" in name:
        return "platform"
    if "报价单位" in sn or "询价单位" in sn:
        return "device"
    return "unknown"


# ---- sheet classification ---------------------------------------------------
OPS_HEADER_HINTS = {"token", "调用", "运维", "运营", "活跃系数", "年调用"}
# deploy 表靠"表名含部署/交付"或"独有列头(成本属性/成本大类/现场实施)"识别；
# 不再用"部署/实施"等宽泛子串匹配表头——否则人天阶段列"实施部署"会误判整表为 deploy。
DEPLOY_NAME_HINTS = {"部署", "交付"}
DEPLOY_HEADER_HINTS = {"成本属性", "成本大类", "现场实施"}
ITEMS_HEADER_MARKERS = {"建设内容", "建设详情", "详情", "模块", "子模块", "智能体", "模型"}
SUMMARY_HEADER_HINTS = {"汇总", "总报价金额", "分项"}


DEVICE_HEADER_HINTS = {"设备名称", "产品名称", "型号", "规格", "参考品牌", "品牌"}


def classify_sheet(name: str, header: list[str | None]) -> str:
    hs = " ".join([str(h) for h in header if h])
    if any(k in name for k in ["汇总", "总表"]) or any(k in hs for k in SUMMARY_HEADER_HINTS):
        return "summary"
    if any(k in hs for k in OPS_HEADER_HINTS):
        return "ops"
    # deploy：表名含部署/交付，或有 deploy 专有列头；且不是明细表（明细表带 建设内容/模块 等标记列）
    if (any(k in name for k in DEPLOY_NAME_HINTS) or any(k in hs for k in DEPLOY_HEADER_HINTS)) \
            and not any(k in hs for k in ITEMS_HEADER_MARKERS):
        return "deploy"
    # device sheets use 设备名称/产品名称 + 数量 + 单价 + 金额 schema
    if any(k in hs for k in DEVICE_HEADER_HINTS):
        return "items"
    if any(k in hs for k in {"建设内容", "建设详情", "详情", "模块", "智能体", "模型"}):
        return "items"
    return "unknown"


# ---- header detection -------------------------------------------------------
HEADER_HINTS = {"序号", "建设内容", "详情", "模块", "单价", "成本", "人/天", "人天", "数量"}


def detect_header_row(ws) -> int:
    """Scan first ~10 rows for the one with most header-like cells."""
    best, best_score = 1, -1
    for r in range(1, min(ws.max_row + 1, 11)):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        score = sum(1 for v in vals if isinstance(v, str) and any(h in v for h in HEADER_HINTS))
        if score > best_score:
            best, best_score = r, score
    return best if best_score > 0 else 1


def parse_workbook(path: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets_out = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        if ws.max_row is None or ws.max_row < 1:
            continue
        hr = detect_header_row(ws)
        header = [ws.cell(row=hr, column=c).value for c in range(1, ws.max_column + 1)]
        kind = classify_sheet(sn, header)

        rows = []
        for r in range(hr + 1, ws.max_row + 1):
            cells_raw = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in cells_raw):
                continue
            row_obj = {
                "row_index": r,
                "cells": {
                    (header[i] if header[i] else f"col_{i+1}"): cells_raw[i]
                    for i in range(len(cells_raw))
                    if cells_raw[i] is not None
                },
            }
            rows.append(row_obj)

        sheets_out.append(
            {
                "name": sn,
                "kind": kind,
                "header_row": hr,
                "columns": [
                    {"index": i + 1, "header": header[i]}
                    for i in range(len(header))
                    if header[i]
                ],
                "rows": rows,
                "merged_cells": [str(rng) for rng in ws.merged_cells.ranges],
                "max_row": ws.max_row,
                "max_col": ws.max_column,
            }
        )

    return {
        "path": str(path),
        "kind": classify_workbook(path.name, wb.sheetnames),
        "sheets": sheets_out,
    }


def main(quote_paths: list[str], out_path: str) -> None:
    workbooks = [parse_workbook(Path(p)) for p in quote_paths]
    result = {"workbooks": workbooks, "schema_version": 1}
    Path(out_path).write_text(json.dumps(result, ensure_ascii=False, indent=2))
    n_sheets = sum(len(w["sheets"]) for w in workbooks)
    n_rows = sum(len(s["rows"]) for w in workbooks for s in w["sheets"])
    print(f"parsed {len(workbooks)} workbooks, {n_sheets} sheets, {n_rows} rows → {out_path}")


if __name__ == "__main__":  # pragma: no cover
    if len(sys.argv) < 3:
        sys.exit("usage: python parse_quote.py <out.json> <xlsx> [<xlsx>...]")
    main(sys.argv[2:], sys.argv[1])
